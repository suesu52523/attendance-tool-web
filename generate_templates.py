import os
from datetime import date, time
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
from openpyxl.utils import get_column_letter

TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), 'templates')
os.makedirs(TEMPLATES_DIR, exist_ok=True)

# 通用样式
header_font = Font(bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='0071E3', end_color='0071E3', fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
thin_border = Border(
    left=Side(style='thin', color='E5E5E5'),
    right=Side(style='thin', color='E5E5E5'),
    top=Side(style='thin', color='E5E5E5'),
    bottom=Side(style='thin', color='E5E5E5'),
)


def style_header(ws, headers):
    """为表头设置样式和列宽"""
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, len(str(header)) * 2.2)
    ws.row_dimensions[1].height = 28


def save_template(wb, filename):
    path = os.path.join(TEMPLATES_DIR, filename)
    wb.save(path)
    print(f'Generated: {path}')


# 班组列表（允许新增，这里是默认示例）
GROUPS = [
    '底盘一组', '底盘二组', '前悬一组', '前悬二组',
    '车门一组', '车门二组', '电装一组', '电装二组',
    '复合一组', '复合二组', '总装一组', '总装二组',
    '总装三组', '总装四组', '总装五组', '总装六组',
]


def create_group_template():
    """班组填报表模板：只保留核心字段，增加起止日期/时间，工时自动计算"""
    wb = Workbook()
    headers = [
        '序号', '工号', '姓名', '班组',
        '加班开始日期', '加班开始时间', '加班结束日期', '加班结束时间', '加班时数',
        '加班原因', '加班类别', '科负责人核准'
    ]

    first = True
    for group in GROUPS:
        if first:
            ws = wb.active
            ws.title = group
            first = False
        else:
            ws = wb.create_sheet(title=group)
        style_header(ws, headers)

        # 示例行：日期与时间使用真实 Excel 日期/时间类型，便于公式计算
        ws.append([
            1, '10010001', '张三', group,
            date(2026, 8, 1), time(15, 45), date(2026, 8, 1), time(17, 35), None,
            '产能爬坡', '工作日', '核准'
        ])
        ws.append([
            2, '10010002', '李四', group,
            date(2026, 8, 2), time(7, 0), date(2026, 8, 2), time(15, 45), None,
            '设备检修', '休息日', '核准'
        ])

        # 设置日期/时间列格式
        for row in [2, 3]:
            ws.cell(row=row, column=5).number_format = 'yyyy-mm-dd'  # 加班开始日期
            ws.cell(row=row, column=6).number_format = 'h:mm'        # 加班开始时间
            ws.cell(row=row, column=7).number_format = 'yyyy-mm-dd'  # 加班结束日期
            ws.cell(row=row, column=8).number_format = 'h:mm'        # 加班结束时间
            # 加班时数公式：((结束日期+结束时间)-(开始日期+开始时间))*24
            hour_cell = ws.cell(row=row, column=9, value=f'=IF(OR(ISBLANK(E{row}),ISBLANK(F{row}),ISBLANK(G{row}),ISBLANK(H{row})),"",((G{row}+H{row})-(E{row}+F{row}))*24)')
            hour_cell.number_format = '0.00'
            hour_cell.border = thin_border

    save_template(wb, '班组填报表模板.xlsx')


def create_abnormal_template():
    """异常表模板：基于加班考勤校对(40)字段，含 T-1/T0/T+1 三日排班与转换区间、就餐信息"""
    wb = Workbook()
    ws = wb.active
    ws.title = '异常记录'
    headers = [
        'ID', '工号', '姓名', '科室',
        'T-1日*系统排班', 'T0日*系统排班', 'T+1日*系统排班',
        '开始加班打卡时间', '结束加班打卡时间',
        '开始日期', '结束日期', '开始时间', '结束时间', '转换加班区间',
        '上报加班时数', '实际加班时数(未减吃饭时间)', '差异',
        '提醒信息', '就餐信息'
    ]
    style_header(ws, headers)
    ws.append([
        1, '10010001', '张三', '底盘一组',
        '双班中班 2026-07-31 15:45:00~2026-08-01 00:20:00',
        '双班早班 2026-08-01 07:00:00~2026-08-01 15:45:00',
        '固定班 2026-08-02 08:45:00~2026-08-02 17:30:00',
        '', '2026-08-01 16:56:29',
        '20260801', '20260801', '15:45', '18:45', '15:45-18:45',
        3, 1, 2,
        '【加班结束卡】请确认加班结束时间', '无'
    ])
    save_template(wb, '异常表模板.xlsx')


def create_rectify_template():
    """整改表模板：按班组拆分为多个 sheet，字段与异常表一致并补充处置列"""
    wb = Workbook()
    base_headers = [
        'ID', '工号', '姓名', '科室',
        'T-1日*系统排班', 'T0日*系统排班', 'T+1日*系统排班',
        '开始加班打卡时间', '结束加班打卡时间',
        '开始日期', '结束日期', '开始时间', '结束时间', '转换加班区间',
        '上报加班时数', '实际加班时数(未减吃饭时间)', '差异',
        '提醒信息', '就餐信息'
    ]
    action_headers = [
        '处置方式', '修改后开始日期', '修改后开始时间', '修改后结束日期', '修改后结束时间',
        '修改后上报加班时数', '调班日期', '调班班次', '异常说明（必填）'
    ]
    headers = base_headers + action_headers

    first = True
    for group in GROUPS[:4]:
        if first:
            ws = wb.active
            ws.title = group
            first = False
        else:
            ws = wb.create_sheet(title=group)
        style_header(ws, headers)
        ws.append([
            1, '10010001', '张三', group,
            '双班中班 2026-07-31 15:45:00~2026-08-01 00:20:00',
            '双班早班 2026-08-01 07:00:00~2026-08-01 15:45:00',
            '固定班 2026-08-02 08:45:00~2026-08-02 17:30:00',
            '', '2026-08-01 16:56:29',
            '20260801', '20260801', '15:45', '18:45', '15:45-18:45',
            3, 1, 2,
            '【加班结束卡】请确认加班结束时间', '无',
            '修改', '20260801', '18:00', '20260801', '21:00', 3, '', '', ''
        ])

    save_template(wb, '整改表模板.xlsx')


def create_system_template():
    """系统数据模板：严格匹配上传系统模板的数据.xls 的 2007-加班申请 前 6 行表头结构"""
    wb = Workbook()
    ws = wb.active
    ws.title = '2007-加班申请'

    # 11 列字段，不改变字段数量
    headers = [
        '中文名称', '工号', '姓名', '开始日期', '结束日期',
        '类型', '开始时间', '结束时间', '定额量', '加班报酬类型', '加班原因'
    ]

    # 前 6 行：中文名称、选项值、是否必填、填报说明、取值举例、序列号
    row1 = headers  # 中文名称
    row2 = ['选项值', '', '', '', '', 'T556P', '', '', '', 'T555R', '']
    row3 = ['是否必填', '', '', 'X', 'X', 'X', 'X', 'X', 'X', 'X', 'X']
    row4 = ['填报说明', '使用新方案的8位数字工号',
            '此字段作为校验字段，如与个人基本信息内姓名不一致，导入会报错',
            'YYYYMMDD格式', 'YYYYMMDD格式', '', 'HHMM格式', 'HHMM格式',
            '时数', '选项', '自由填写']
    row5 = ['取值举例', '163161', '谭文杰', '20200101', '20200101',
            '', '0845', '1730', '5', '3 补休', '测试']
    row6 = ['序列号', 'RP50G-PERNR', '', 'P2007-BEGDA', 'P2007-ENDDA',
            'P2007-KTART', 'P2007-BEGUZ', 'P2007-ENDUZ',
            'P2007-ANZHL', 'P2007-VERSL', 'RESON']

    for row in [row1, row2, row3, row4, row5, row6]:
        ws.append(row)

    # 示例数据行（第 7 行起）
    ws.append([
        1, '00151177', '童辉武', '20260701', '20260701',
        '10 已核准的加班', '15:45', '18:45', 3.0, '1 支付加班费', '生产一线早班统一加班生产'
    ])
    ws.append([
        2, '00151177', '童辉武', '20260702', '20260702',
        '10 已核准的加班', '15:45', '18:45', 3.0, '1 支付加班费', '生产一线早班统一加班生产'
    ])

    # 码表
    code_ws = wb.create_sheet(title='码表')
    code_ws.append(['类型码', '加班报酬类型码'])
    code_ws.append(['10 已核准的加班', '1 支付加班费'])
    code_ws.append(['', '3 补休'])

    save_template(wb, '系统数据模板.xlsx')


def create_shift_template():
    """调班数据导入模板：严格匹配 2003_ZP 副本的多个 sheet 结构"""
    wb = Workbook()

    # 2003-调班信息（6 列主模板）
    ws1 = wb.active
    ws1.title = '2003-调班信息'
    headers1 = ['中文名称', '工号', '姓名', '开始日期', '结束日期', '日工作计划']
    style_header(ws1, headers1)
    ws1.append([1, '10010001', '张三', '20260902', '20260902', 'OFF 休息'])

    # Sheet2（8 列扩展模板）
    ws2 = wb.create_sheet(title='Sheet2')
    headers2 = ['中文名称', '工号', '姓名', '开始日期', '结束日期', '日工作计划', '出勤项目分类', '备注']
    style_header(ws2, headers2)
    ws2.append(['例子', '10010001', '张三', '20260301', '20260301', 'SF17 固定班', '已反馈出勤计划', '8点45-17点30分'])
    ws2.append(['例子', '10010001', '张三', '20260301', '20260301', 'SF04 双班早班', '其他', '参加公司统一安排军人集训'])

    # 上班时间段
    ws3 = wb.create_sheet(title='上班时间段')
    ws3.append(['班次代码', '上班时间', '备注'])
    ws3.append(['SF04 双班早班', '7:00-15:45', ''])
    ws3.append(['SF17 固定班', '8:45-17:30', ''])
    ws3.append(['SF05 双班中班', '15:45-00:20', ''])
    ws3.append(['SF10 二线中班1545', '15:45-00:30', ''])
    ws3.append(['SF11 二线中班1645', '16:45-1:30', ''])
    ws3.append(['SF12 二线中班1755', '17:55-2:40', '2:40-3:55'])
    ws3.append(['SF13 二线中班1845', '18:45-3:45', ''])

    # 名单
    ws4 = wb.create_sheet(title='名单')
    ws4.append(['工号', '姓名'])
    ws4.append(['10010001', '张三'])
    ws4.append(['10010002', '李四'])

    # Sheet1（空表占位）
    wb.create_sheet(title='Sheet1')

    # 码表（隐藏）
    code_ws = wb.create_sheet(title='码表')
    code_ws.append(['班次代码'])
    for code in ['NS 未排班', 'OFF 休息', 'SF01 正常班', 'SF02 动总二厂正常班',
                 'SF04 双班早班', 'SF05 双班中班', 'SF17 固定班',
                 'SF28 一部改造期早班', 'SF29 一部改造期中班', 'X260 自定义班次1']:
        code_ws.append([code])
    code_ws.sheet_state = 'hidden'

    save_template(wb, '调班数据导入模板.xlsx')


if __name__ == '__main__':
    create_group_template()
    create_abnormal_template()
    create_rectify_template()
    create_system_template()
    create_shift_template()
    print('All templates generated successfully.')
